import base64
import builtins
import csv
import datetime
import io
import json
import re
import shutil
import time

from PIL import Image, ImageOps

struct = wiz.model("struct")


FIELDS = [
    "name", "company", "department", "position", "email", "mobile", "phone",
    "address", "website", "front_image", "back_image", "tags", "memo", "source",
]
IMPORT_FIELDS = [
    "name", "company", "department", "position", "email", "mobile", "phone",
    "address", "website", "tags", "memo", "source",
]
IMPORT_FIELD_OPTIONS = [
    dict(key="", label="가져오지 않음"),
    dict(key="name", label="이름"),
    dict(key="company", label="회사"),
    dict(key="department", label="부서"),
    dict(key="position", label="직책"),
    dict(key="email", label="이메일"),
    dict(key="mobile", label="휴대폰"),
    dict(key="phone", label="전화"),
    dict(key="address", label="주소"),
    dict(key="website", label="웹사이트"),
    dict(key="memo", label="메모"),
    dict(key="tags", label="태그"),
    dict(key="source", label="출처"),
]
EXPORT_COLUMNS = [
    ("name", "이름"),
    ("company", "회사"),
    ("department", "부서"),
    ("position", "직책"),
    ("email", "이메일"),
    ("mobile", "휴대폰"),
    ("phone", "전화"),
    ("address", "주소"),
    ("website", "웹사이트"),
    ("memo", "메모"),
    ("tags", "태그"),
    ("source", "출처"),
    ("created", "생성일"),
    ("updated", "수정일"),
]
IMPORT_ALIASES = {
    "name": ["name", "full name", "person", "contact", "성명", "이름", "담당자", "담당자명"],
    "company": ["company", "company name", "organization", "office", "회사", "회사명", "소속", "기관", "기관명"],
    "department": ["department", "dept", "division", "team", "부서", "부서명", "팀", "본부"],
    "position": ["position", "title", "role", "job title", "직책", "직위", "역할", "직급"],
    "email": ["email", "e-mail", "mail", "이메일", "메일"],
    "mobile": ["mobile", "cell", "cellphone", "phone mobile", "휴대폰", "핸드폰", "모바일", "휴대전화"],
    "phone": ["phone", "tel", "telephone", "office phone", "전화", "전화번호", "대표번호", "유선"],
    "address": ["address", "addr", "location", "주소", "소재지", "위치"],
    "website": ["website", "web", "homepage", "url", "site", "웹사이트", "홈페이지", "사이트"],
    "memo": ["memo", "note", "notes", "remark", "remarks", "메모", "비고", "노트"],
    "tags": ["tag", "tags", "label", "labels", "태그", "라벨", "분류"],
    "source": ["source", "origin", "출처", "소스"],
}
IMPORT_EXTENSIONS = [".csv", ".txt", ".xlsx"]
IMPORT_DELIMITERS = {
    "comma": ",",
    "tab": "\t",
    "semicolon": ";",
    "pipe": "|",
}
MAX_IMPORT_ROWS = 5000
MAX_EXPORT_ROWS = 10000
MAX_IMPORT_BYTES = 6 * 1024 * 1024

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?:\+\d{1,3}[\s.-]?\d{1,3}[\s.-]?\d{3,4}[\s.-]?\d{4}|0\d{1,2}[\s.-]?\d{3,4}[\s.-]?\d{4}|\d{3}[\s.-]\d{3,4}[\s.-]\d{4})")
URL_RE = re.compile(r"(?<!@)\b(?:https?://)?(?:www\.)?[A-Z0-9-]+(?:\.[A-Z0-9-]+)*\.[A-Z]{2,}(?:/[\w.?#=&%-]*)?", re.I)
COMPANY_RE = re.compile(
    r"주식회사\s*[가-힣A-Za-z0-9]+|\(주\)\s*[가-힣A-Za-z0-9]+|㈜\s*[가-힣A-Za-z0-9]+|"
    r"유한회사\s*[가-힣A-Za-z0-9]+|합자회사\s*[가-힣A-Za-z0-9]+|합명회사\s*[가-힣A-Za-z0-9]+|"
    r"재단법인\s*[가-힣A-Za-z0-9]+|사단법인\s*[가-힣A-Za-z0-9]+|"
    r"[가-힣A-Za-z0-9]+(?:테크|랩스|연구소|스튜디오|파트너스|솔루션|시스템즈|컴퍼니)(?:\s*(?:연구소|랩스|스튜디오))?|"
    r"[A-Z][A-Za-z0-9&.,\s-]{1,80}?\b(?:Co\.?\s*Ltd\.?|Ltd\.?|Limited|LLC|LLP|Inc\.?|Corp\.?|Corporation|Company|GmbH|Pte\.?\s*Ltd\.?|Pty\.?\s*Ltd\.?|PLC|Group|Labs?|Studio|Partners|Technologies|Systems)",
    re.I,
)
COMPANY_HINT_RE = re.compile(
    r"회사|그룹|컴퍼니|코퍼레이션|법인|재단|사단|연구소|스튜디오|랩스?|파트너스|솔루션|테크|시스템즈|"
    r"\b(?:Group|Holdings|Labs?|Studio|Studios|Partners|Solutions|Technologies|Technology|Tech|Systems|Services|Global)\b",
    re.I,
)
POSITION_RE = re.compile(
    r"대표이사|공동대표|대표|회장|부회장|사장|부사장|전무|상무|이사|감사|고문|실장|팀장|본부장|센터장|"
    r"부장|차장|과장|대리|주임|사원|수석매니저|선임매니저|책임매니저|프로덕트매니저|프로젝트매니저|책임|수석|선임|연구원|매니저|컨설턴트|디자이너|개발자|엔지니어|"
    r"Co-Founder|Founder|Owner|President|Vice President|VP|Managing Director|General Manager|"
    r"Chief\s+[A-Za-z\s]+Officer|CEO|CTO|CFO|COO|CPO|CMO|CIO|Principal|Partner|"
    r"Senior|Junior|Staff|Associate|Consultant|Specialist|Designer|Developer|Researcher|Analyst|Architect|"
    r"Product\s+Manager|Project\s+Manager|Sales\s+Manager|Marketing\s+Manager|Manager|Director|Lead|Head|Engineer|Officer",
    re.I,
)
ROLE_RE = re.compile(
    POSITION_RE.pattern
    + r"|본부|센터|부서|팀|사업부|부문|실|국|과|파트|랩|연구소|"
    + r"Team|Dept\.?|Department|Division|Office|Center|Centre|Lab|Laboratory|Group|Unit|Part|Squad|Chapter|"
    + r"Software|Product|Design|Planning|Strategy|Sales|Marketing|Operation|Operations|Engineering|Development",
    re.I,
)
ADDRESS_RE = re.compile(
    r"\bAddr\.?|\bAddress\b|주소|소재지|위치|자치시|특별시|광역시|시\s|군\s|구\s|동\s|읍\s|면\s|"
    r"대로|로\s|길\s|번지|호\s|층\s|빌딩|타워|프라자|센터|Republic of Korea|Korea|"
    r"Hannuri|Hanlim|Suite|\bSte\.?\b|Floor|\bFl\.?\b|\bBldg\.?\b|Building|Road|\bRd\.?\b|Street|\bSt\.?\b|\bAve\.?\b|Avenue|\bBlvd\.?\b|Drive|\bDr\.?\b|\bCity\b|\bDong\b|\bGu\b|\bSi\b",
    re.I,
)
ADDRESS_STRONG_RE = re.compile(
    r"\bAddr\.?|\bAddress\b|\bHQ\b|Headquarters|Location|Located at|주소|소재지|위치|자치시|특별시|광역시|"
    r"시\s|군\s|구\s|동\s|읍\s|면\s|대로|로\s|길\s|번지|호\s|층\s|빌딩|타워|프라자|Republic of Korea|Korea|"
    r"Hannuri|Hanlim|Suite|\bSte\.?\b|Floor|\bFl\.?\b|\bBldg\.?\b|Building|Road|\bRd\.?\b|Street|\bSt\.?\b|\bAve\.?\b|Avenue|\bBlvd\.?\b|Drive|\bDr\.?\b|\bCity\b|\bDong\b|\bGu\b|\bSi\b",
    re.I,
)
NAME_LABEL_RE = re.compile(r"^(?:성명|이름|담당자|Name|Contact|Person in Charge)\s*[:：.-]?\s*(.+)$", re.I)
MOBILE_LABEL_RE = re.compile(r"\bM(?:obile)?\.?\b|\bMob\.?\b|\bCell(?:phone|ular)?\.?\b|\bC\.?P\.?\b|\bH\.?P\.?\b|휴대폰|핸드폰|모바일|연락처|휴대전화|휴대|무선", re.I)
PHONE_LABEL_RE = re.compile(r"\bT(?:el)?\.?\b|\bTelephone\.?\b|\bPhone\.?\b|\bOffice\.?\b|\bDirect\.?\b|\bMain\.?\b|전화번호|전화|대표전화|대표번호|유선|사무실|내선", re.I)
FAX_LABEL_RE = re.compile(r"\bF(?:ax)?\.?\b|팩스", re.I)
OCR_CONFIGS = ["--psm 11", "--psm 6"]
OCR_TIMEOUT_SECONDS = 25
MERGE_FIELDS = ["name", "company", "department", "position", "email", "mobile", "phone", "address", "website"]


def _current_user_id():
    user_id = wiz.session.get("id", "")
    if not user_id:
        wiz.response.status(401, message="로그인이 필요합니다.")
    return user_id


def _payload():
    data = {}
    for key in FIELDS:
        data[key] = wiz.request.query(key, "")
    return data


def _clean_line(line):
    return re.sub(r"\s+", " ", re.sub(r"[|•]", " ", line or "")).strip(" -_:;,.")


def _strip_contact(value):
    value = EMAIL_RE.sub("", value or "")
    value = URL_RE.sub("", value)
    value = PHONE_RE.sub("", value)
    value = re.sub(
        r"\b(?:M|T|F|Tel|Telephone|Phone|Mobile|Mob|Cell|Fax|Direct|Office|Main|E-mail|Email|Mail|Web|Website|Homepage|Addr|Address)\.?\b|"
        r"휴대폰|핸드폰|모바일|휴대전화|전화번호|대표전화|대표번호|전화|팩스|이메일|메일|홈페이지|웹사이트|주소|소재지",
        " ",
        value,
        flags=re.I,
    )
    return _clean_line(value)


def _clean_company(line):
    match = COMPANY_RE.search(line or "")
    if match is not None:
        return _strip_contact(match.group(0))

    candidate = _strip_contact(line or "")
    if (
        candidate
        and COMPANY_HINT_RE.search(candidate)
        and ROLE_RE.search(candidate) is None
        and ADDRESS_STRONG_RE.search(candidate) is None
        and not _is_contact_line(candidate)
    ):
        return candidate
    return ""


def _extract_websites(text):
    email_spans = [match.span() for match in EMAIL_RE.finditer(text or "")]
    websites = []
    for match in URL_RE.finditer(text or ""):
        start, end = match.span()
        if any(start >= email_start and end <= email_end for email_start, email_end in email_spans):
            continue
        websites.append(match.group(0))
    return websites


def _is_contact_line(line):
    return EMAIL_RE.search(line or "") or URL_RE.search(line or "") or PHONE_RE.search(line or "")


def _normalize_english_name(value):
    words = []
    for item in re.split(r"\s+", value or ""):
        if not item:
            continue
        word = item[:1].upper() + item[1:].lower()
        if word.lower() in ["kwon", "kwonn"]:
            word = "Kwon"
        words.append(word)
    return " ".join(words)


def _korean_name_from(line):
    match = re.search(
        r"((?:[가-힣]\s*){2,4})(?=\s*(?:개발|영업|기획|마케팅|전략|디자인|사업|팀|본부|센터|부서|사업부|부문|실|국|과|파트|랩|연구소|대표|이사|실장|팀장|부장|차장|과장|대리|주임|사원|책임|수석|선임|연구원|매니저|컨설턴트|디자이너|개발자|엔지니어|$))",
        line or "",
    )
    if match is None:
        return ""
    name = re.sub(r"\s+", "", match.group(1))
    if re.search(r"개발|영업|기획|마케팅|전략|디자인|사업|팀|본부|센터|부서|실|국|과|파트|랩|연구|책임|수석|선임|매니저", name):
        return ""
    if 2 <= len(name) <= 4:
        return name
    return ""


def _english_name_from(line):
    match = re.search(
        r"\b([A-Z][A-Za-z]{1,24})\s+([A-Z][A-Za-z]{1,24})(?=\s+(?:Co-Founder|Founder|Owner|President|Vice|VP|Managing|General|Chief|CEO|CTO|CFO|COO|CPO|CMO|CIO|Principal|Partner|Senior|Junior|Staff|Associate|Consultant|Specialist|Designer|Developer|Researcher|Analyst|Architect|Product|Project|Sales|Marketing|Lead|Head|Manager|Director|Engineer|Officer|of\b))",
        line or "",
    )
    if match is None:
        return ""
    name = _normalize_english_name(match.group(0))
    if re.search(r"Chief|Product|Project|Sales|Marketing|Senior|Junior|Staff|Associate|Consultant|Specialist|Designer|Developer|Researcher|Analyst|Architect|Manager|Director|Engineer|Officer|Founder|President|Lead|Head", name, re.I):
        return ""
    return name


def _extract_name(lines):
    for line in lines:
        match = NAME_LABEL_RE.search(line)
        if match and COMPANY_RE.search(line) is None:
            candidate = _strip_contact(match.group(1))
            korean = _korean_name_from(candidate)
            if korean:
                return korean
            english = re.match(r"^([A-Z][A-Za-z]{1,24}(?:\s+[A-Z][A-Za-z]{1,24}){1,2})$", candidate)
            if english:
                return _normalize_english_name(english.group(1))

    for line in lines:
        if COMPANY_RE.search(line) or _is_contact_line(line) or ADDRESS_STRONG_RE.search(line):
            continue
        if ROLE_RE.search(line):
            name = _korean_name_from(line) or _english_name_from(line)
            if name:
                return name

    for line in lines:
        if _is_contact_line(line) or ADDRESS_STRONG_RE.search(line) or COMPANY_RE.search(line):
            continue
        name = _korean_name_from(line)
        if name and len(line.replace(" ", "")) <= len(name) + 2:
            return name
        match = re.match(r"^([A-Z][A-Za-z]{1,24}\s+[A-Z][A-Za-z]{1,24})$", line)
        if match:
            return _normalize_english_name(match.group(1))
    return ""


def _extract_role(lines, name):
    department = ""
    position = ""
    name_variants = []
    if name:
        name_variants.append(name)
        if re.search(r"[가-힣]", name):
            name_variants.append(" ".join(name))

    for line in lines:
        if ROLE_RE.search(line) is None:
            continue
        if COMPANY_RE.search(line) or _is_contact_line(line) or ADDRESS_STRONG_RE.search(line):
            continue

        work = line
        for variant in name_variants:
            work = re.sub(re.escape(variant), " ", work, flags=re.I)

        if not position:
            match = POSITION_RE.search(work)
            if match:
                position = match.group(0)

        if not department:
            match = re.search(
                r"(?:Lead|Head|Manager|Director|Engineer|Officer|Consultant|Specialist)\s+(?:of|for)\s+"
                r"([A-Z][A-Za-z0-9\s&/-]{2,70}?\s+(?:Team|Dept\.?|Department|Division|Office|Center|Centre|Lab|Laboratory|Group|Unit|Part|Squad|Chapter))\b",
                work,
                re.I,
            )
            if match:
                department = _clean_line(match.group(1))
            else:
                match = re.search(
                    r"\b([A-Z][A-Za-z0-9\s&/-]{2,70}?\s+(?:Team|Dept\.?|Department|Division|Office|Center|Centre|Lab|Laboratory|Group|Unit|Part|Squad|Chapter))\b",
                    work,
                    re.I,
                )
                if match:
                    department = _clean_line(match.group(1))
                else:
                    normalized = re.sub(r"[^가-힣A-Za-z0-9\s/]", " ", work)
                    match = re.search(r"([가-힣A-Za-z0-9]{1,28}(?:팀|본부|센터|부서|사업부|부문|실|국|과|파트|랩|연구소))", normalized)
                    if match:
                        department = match.group(1)

        if department and position:
            break

    return department, position


def _normalize_phone(value):
    value = _clean_line(value)
    value = re.sub(r"(?<=\d)\s+(?=\d)", "-", value)
    return re.sub(r"\s*[-.]\s*", "-", value)


def _extract_phone_fields(lines):
    mobile = ""
    phone = ""
    fallback_numbers = []

    for line in lines:
        numbers = [_normalize_phone(item) for item in PHONE_RE.findall(line)]
        if not numbers:
            continue

        is_fax = FAX_LABEL_RE.search(line) is not None and MOBILE_LABEL_RE.search(line) is None and PHONE_LABEL_RE.search(line) is None
        is_mobile = MOBILE_LABEL_RE.search(line) is not None
        is_phone = PHONE_LABEL_RE.search(line) is not None

        for number in numbers:
            normalized = re.sub(r"\D", "", number)
            looks_mobile = normalized.startswith("01") or normalized.startswith("8210")
            if is_fax:
                continue
            if looks_mobile:
                if not mobile:
                    mobile = number
                continue
            if not mobile and is_mobile:
                mobile = number
            elif not phone and is_phone:
                phone = number
            else:
                fallback_numbers.append(number)

    for number in fallback_numbers:
        normalized = re.sub(r"\D", "", number)
        if not mobile and (normalized.startswith("01") or normalized.startswith("8210")):
            mobile = number
        elif not phone and number != mobile:
            phone = number

    return mobile, phone


def _extract_address(lines):
    candidates = []
    for index, line in enumerate(lines):
        if ADDRESS_RE.search(line) is None:
            continue
        for span in [1, 2, 3]:
            candidate = " ".join(lines[index:index + span])
            candidate = re.sub(r"^.*?\b(?:Addr|Address|HQ|Headquarters|Office|Location|Located at)[,.]?\s*", "", candidate, flags=re.I)
            candidate = re.sub(r"^(?:주소|소재지|사업장|사무실|본사|지점)[:：.\s]*", "", candidate)
            candidate = COMPANY_RE.split(candidate, 1)[0]
            candidate = re.split(r"\b(?:주식회사|Season\s+Co|EASA|pee)\b", candidate, 1, flags=re.I)[0]
            candidate = _strip_contact(candidate)
            candidate = re.sub(r"\b(?:om|ZSy|aN|ee|oe|se|Ly|200)\b", " ", candidate, flags=re.I)
            candidate = _clean_line(candidate)
            if len(candidate) >= 8:
                candidates.append(candidate)

    if not candidates:
        return ""

    def score(candidate):
        keyword_score = sum(bool(re.search(pattern, candidate, re.I)) for pattern in [
            r"자치시|특별시|광역시|Sejong|Seoul|Busan|\bCity\b",
            r"대로|로\s|길\s|daero|Road|\bRd\.?\b|Street|\bSt\.?\b|\bAve\.?\b|Avenue",
            r"프라자|빌딩|타워|센터|plaza|\bBldg\.?\b|Building|Suite|Floor",
            r"Republic|Korea|USA",
            r"한림|Hanlim",
        ])
        return keyword_score * 20 + len(candidate)

    return sorted(candidates, key=score, reverse=True)[0]


def _merge_ocr_texts(texts):
    lines = []
    seen = set()
    for text in texts:
        for line in re.split(r"[\n\r]+", text or ""):
            cleaned = _clean_line(line)
            if not cleaned:
                continue
            key = cleaned.lower()
            if key in seen:
                continue
            seen.add(key)
            lines.append(cleaned)
    return "\n".join(lines)


def _score_ocr_text(text):
    raw = text or ""
    patterns = [EMAIL_RE, PHONE_RE, URL_RE, COMPANY_RE, ROLE_RE, ADDRESS_RE]
    return sum(1 for pattern in patterns if pattern.search(raw)) + min(len(raw.strip()) // 120, 3)


def _confidence_from_data(data):
    confidences = []
    for value in data.get("conf", []):
        try:
            score = float(value)
            if score >= 0:
                confidences.append(score)
        except Exception:
            pass
    return round(sum(confidences) / len(confidences)) if confidences else 0


def _estimate_confidence(fields, text):
    score = sum(1 for key in MERGE_FIELDS if fields.get(key))
    if fields.get("mobile") or fields.get("phone"):
        score += 1
    if fields.get("email"):
        score += 1
    text_bonus = min(len((text or "").strip()) // 240, 4)
    return min(96, 18 + (score * 7) + text_bonus)


def _field_score(fields):
    score = 0
    for key in ["name", "company", "email", "mobile", "phone"]:
        if fields.get(key):
            score += 3
    for key in ["department", "position", "address", "website"]:
        if fields.get(key):
            score += 1
    return score


def _needs_ocr_boost(fields):
    if not fields.get("name"):
        return True
    if not fields.get("company"):
        return True
    if not fields.get("email") and not fields.get("mobile") and not fields.get("phone"):
        return True
    if not fields.get("department") or not fields.get("position"):
        return True
    if len(fields.get("address", "")) < 12:
        return True
    return False


def _needs_rotation_probe(fields, text):
    if _field_score(fields) >= 6:
        return False
    has_anchor = any(fields.get(key) for key in ["company", "email", "mobile", "phone", "website"])
    if not has_anchor and _score_ocr_text(text) <= 4:
        return True
    if _score_ocr_text(text) <= 1:
        return True
    return len((text or "").strip()) < 40


def _ocr_candidate(pytesseract, prepared, angle, config):
    image = prepared if angle == 0 else prepared.rotate(angle, expand=True)
    try:
        text = pytesseract.image_to_string(image, lang="kor+eng", config=config, timeout=OCR_TIMEOUT_SECONDS)
    except Exception as e:
        return dict(angle=angle, config=config, image=image, text="", fields={}, score=0, error=str(e))
    fields = _parse_ocr_text(text)
    score = (_score_ocr_text(text) * 20) + (_field_score(fields) * 10) + min(len((text or "").strip()) // 30, 10)
    return dict(angle=angle, config=config, image=image, text=text, fields=fields, score=score)


def _merge_ocr_fields(field_sets):
    merged = {key: "" for key in MERGE_FIELDS}
    for fields in field_sets:
        for key in MERGE_FIELDS:
            if not merged.get(key) and fields.get(key):
                merged[key] = fields.get(key)
    return merged


def _parse_ocr_text(text):
    raw = text or ""
    lines = [_clean_line(line) for line in re.split(r"[\n\r]+", raw)]
    lines = [line for line in lines if line]
    joined = " ".join(lines)

    email_match = EMAIL_RE.search(joined)
    website_matches = _extract_websites(joined)

    company = ""
    for line in lines:
        company = _clean_company(line)
        if company:
            break

    name = _extract_name(lines)
    department, position = _extract_role(lines, name)

    mobile, phone = _extract_phone_fields(lines)

    return dict(
        name=name,
        company=company,
        department=department,
        position=position,
        email=email_match.group(0) if email_match else "",
        mobile=mobile,
        phone=phone,
        address=_extract_address(lines),
        website=website_matches[0] if website_matches else "",
    )


def _decode_image(image_data):
    if not image_data:
        raise Exception("이미지 데이터가 필요합니다.")
    if image_data.startswith("data:") and "," in image_data:
        image_data = image_data.split(",", 1)[1]
    raw = base64.b64decode(image_data)
    if len(raw) > 8 * 1024 * 1024:
        raise Exception("이미지 용량이 너무 큽니다.")
    image = Image.open(io.BytesIO(raw))
    return ImageOps.exif_transpose(image)


def _prepare_image(image):
    image = image.convert("RGB")
    image.thumbnail((2200, 2200))
    image = ImageOps.grayscale(image)
    image = ImageOps.autocontrast(image)
    return image


def _server_ocr(image):
    if shutil.which("tesseract") is None:
        return dict(available=False, text="", confidence=0, message="서버 OCR 엔진 없음")

    try:
        import pytesseract
    except Exception:
        return dict(available=False, text="", confidence=0, message="서버 OCR 래퍼 없음")

    try:
        started = time.perf_counter()
        prepared = _prepare_image(image)
        results = []

        first_config = OCR_CONFIGS[0]
        first = _ocr_candidate(pytesseract, prepared, 0, first_config)
        results.append(first)
        selected = first

        if _needs_rotation_probe(first["fields"], first["text"]):
            for angle in [90, 270, 180]:
                results.append(_ocr_candidate(pytesseract, prepared, angle, first_config))
            selected = sorted(results, key=lambda item: item["score"], reverse=True)[0]

        if _needs_ocr_boost(selected["fields"]):
            boost_config = OCR_CONFIGS[1]
            boost = _ocr_candidate(pytesseract, selected["image"], 0, boost_config)
            boost["angle"] = selected["angle"]
            results.append(boost)

        selected_texts = [
            item["text"]
            for item in results
            if item["angle"] == selected["angle"] and item["score"] > 0
        ]
        text = _merge_ocr_texts(selected_texts) or selected["text"]
        errors = [item.get("error") for item in results if item.get("error")]
        if not text and errors:
            return dict(
                available=False,
                text="",
                confidence=0,
                message=errors[0],
                elapsed_ms=round((time.perf_counter() - started) * 1000),
                passes=len(results),
                rotation=selected["angle"],
            )
        fields = _parse_ocr_text(text)
        elapsed_ms = round((time.perf_counter() - started) * 1000)
        return dict(
            available=True,
            text=text,
            confidence=_estimate_confidence(fields, text),
            message="서버 분석 완료",
            elapsed_ms=elapsed_ms,
            passes=len(results),
            rotation=selected["angle"],
        )
    except Exception as e:
        return dict(available=False, text="", confidence=0, message=str(e))


def _ai_ocr(image_data, label):
    try:
        return struct.ai_setting.analyze_image(image_data, label=label)
    except Exception as e:
        return dict(
            available=False,
            text="",
            fields={key: "" for key in MERGE_FIELDS},
            confidence=0,
            quality_notes=[str(e)] if str(e) else [],
            quality_passed=False,
            quality_score=0,
            payload_format="json",
            execution_failed=True,
            message=str(e),
            engine="ai-unavailable",
        )


def _merge_ai_fields(fields, text):
    parsed = _parse_ocr_text(text)
    fields = fields if isinstance(fields, dict) else {}
    merged = {key: "" for key in MERGE_FIELDS}
    for key in MERGE_FIELDS:
        value = fields.get(key, "")
        if isinstance(value, builtins.list):
            value = ", ".join([str(item).strip() for item in value if str(item).strip()])
        if value is None:
            value = ""
        value = _clean_line(str(value)) if key != "address" else _clean_line(str(value))
        merged[key] = value or parsed.get(key, "")

    return merged


def _unique_values(items):
    values = []
    for item in items:
        if item and item not in values:
            values.append(item)
    return values


def _bool_query(name, default=False):
    value = str(wiz.request.query(name, "true" if default else "false")).lower()
    return value in ["1", "true", "yes", "y", "on"]


def _file_extension(filename):
    filename = (filename or "").lower().strip()
    if "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[1]


def _safe_export_name(name):
    cleaned = re.sub(r"[^0-9A-Za-z가-힣._-]+", "-", name or "business-cards").strip(".-")
    return cleaned or "business-cards"


def _decode_upload_data(file_data):
    file_data = file_data or ""
    if "," in file_data and file_data.startswith("data:"):
        file_data = file_data.split(",", 1)[1]
    try:
        raw = base64.b64decode(file_data)
    except Exception:
        raise Exception("파일 데이터를 읽지 못했습니다.")
    if len(raw) > MAX_IMPORT_BYTES:
        raise Exception("가져오기 파일은 6MB 이하만 사용할 수 있습니다.")
    return raw


def _decode_text(raw):
    for encoding in ["utf-8-sig", "utf-8", "cp949", "euc-kr", "utf-16"]:
        try:
            return raw.decode(encoding)
        except Exception:
            continue
    raise Exception("텍스트 인코딩을 확인할 수 없습니다. UTF-8 또는 CP949 파일로 저장해 주세요.")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _column_letter(index):
    index += 1
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _normalize_import_rows(rows):
    normalized = []
    width = 0
    for row in rows:
        values = [_cell_text(value) for value in row]
        while values and not values[-1]:
            values.pop()
        if not any(values):
            continue
        width = max(width, len(values))
        normalized.append(values)
        if len(normalized) > MAX_IMPORT_ROWS + 1:
            break

    for row in normalized:
        if len(row) < width:
            row.extend([""] * (width - len(row)))
    return normalized


def _detect_delimiter(text):
    sample = (text or "")[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        if dialect.delimiter in IMPORT_DELIMITERS.values():
            return dialect.delimiter
    except Exception:
        pass

    scores = []
    for delimiter in IMPORT_DELIMITERS.values():
        scores.append((sample.count(delimiter), delimiter))
    scores.sort(reverse=True)
    return scores[0][1] if scores and scores[0][0] > 0 else ","


def _parse_delimited_import(raw, delimiter_mode):
    text = _decode_text(raw)
    delimiter = IMPORT_DELIMITERS.get(delimiter_mode) or _detect_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _normalize_import_rows(reader), delimiter


def _parse_xlsx_import(raw):
    try:
        import openpyxl
    except Exception:
        raise Exception("XLSX 파일을 읽기 위한 openpyxl 패키지가 필요합니다.")

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = _normalize_import_rows(sheet.iter_rows(values_only=True))
        workbook.close()
        return rows
    except Exception as e:
        raise Exception(f"XLSX 파일을 읽지 못했습니다: {str(e)}")


def _parse_import_file(filename, file_data, delimiter_mode="auto"):
    extension = _file_extension(filename)
    if extension not in IMPORT_EXTENSIONS:
        raise Exception("CSV, TXT, XLSX 파일만 가져올 수 있습니다.")

    raw = _decode_upload_data(file_data)
    if extension == ".xlsx":
        return _parse_xlsx_import(raw), "xlsx"
    rows, delimiter = _parse_delimited_import(raw, delimiter_mode)
    return rows, delimiter


def _normalized_header(value):
    return re.sub(r"[\s_\-./()\\[\\]:：]+", "", str(value or "").strip().lower())


def _guess_field(header, index, has_header):
    normalized = _normalized_header(header)
    if normalized:
        for field, aliases in IMPORT_ALIASES.items():
            for alias in aliases:
                alias_key = _normalized_header(alias)
                if normalized == alias_key or alias_key in normalized:
                    return field

    if not has_header and index < len(IMPORT_FIELDS):
        return IMPORT_FIELDS[index]
    return ""


def _build_import_preview(filename, file_data, delimiter_mode="auto", has_header=True):
    rows, detected_delimiter = _parse_import_file(filename, file_data, delimiter_mode)
    if not rows:
        raise Exception("가져올 데이터가 없습니다.")

    headers = rows[0] if has_header else [f"열 {_column_letter(index)}" for index in range(len(rows[0]))]
    data_rows = rows[1:] if has_header else rows
    width = max(len(headers), max([len(row) for row in data_rows], default=0))
    headers = headers + [""] * (width - len(headers))
    normalized_rows = []
    for row in data_rows[:MAX_IMPORT_ROWS]:
        normalized_rows.append(row + [""] * (width - len(row)))

    columns = []
    mappings = []
    for index in range(width):
        label = headers[index] or f"열 {_column_letter(index)}"
        sample = ""
        for row in normalized_rows:
            if row[index]:
                sample = row[index]
                break
        columns.append(dict(index=index, label=label, sample=sample))
        mappings.append(dict(index=index, header=label, field=_guess_field(label, index, has_header)))

    return dict(
        filename=filename,
        extension=_file_extension(filename),
        detected_delimiter=detected_delimiter,
        has_header=has_header,
        columns=columns,
        mappings=mappings,
        field_options=IMPORT_FIELD_OPTIONS,
        rows=normalized_rows,
        preview=normalized_rows[:12],
        total_rows=len(normalized_rows),
        max_rows=MAX_IMPORT_ROWS,
    )


def _load_mappings(value, columns):
    try:
        raw = json.loads(value or "[]")
    except Exception:
        raw = []

    valid_fields = set(IMPORT_FIELDS)
    mapping = {}
    for item in raw if isinstance(raw, list) else []:
        try:
            index = int(item.get("index", -1))
        except Exception:
            index = -1
        field = item.get("field", "")
        if index >= 0 and field in valid_fields:
            mapping[index] = field

    if not mapping:
        for item in columns:
            field = item.get("field", "")
            if field in valid_fields:
                mapping[int(item.get("index", -1))] = field
    return mapping


def _field_label(field):
    for option in IMPORT_FIELD_OPTIONS:
        if option.get("key") == field:
            return option.get("label") or field
    return field


def _mapping_duplicate_errors(mappings, columns):
    headers = {}
    for column in columns or []:
        try:
            headers[int(column.get("index", -1))] = column.get("label", "")
        except Exception:
            continue

    grouped = {}
    for index, field in mappings.items():
        if not field:
            continue
        grouped.setdefault(field, []).append(headers.get(index) or f"열 {_column_letter(index)}")

    errors = []
    for field, names in grouped.items():
        if len(names) > 1:
            errors.append(dict(field=field, label=_field_label(field), columns=names))
    return errors


def _mapped_import_rows(preview, mappings, append_unmapped_to_memo=False):
    rows = []
    errors = []
    columns = preview.get("columns", [])
    column_headers = {int(column.get("index", -1)): column.get("label", "") for column in columns}
    for row_index, row in enumerate(preview.get("rows", []), start=2 if preview.get("has_header") else 1):
        item = {key: "" for key in IMPORT_FIELDS}
        for index, field in mappings.items():
            if index >= len(row) or not field:
                continue
            value = _cell_text(row[index])
            if value and not item.get(field):
                item[field] = value

        if not any(item.get(key) for key in IMPORT_FIELDS):
            continue
        if not item.get("source"):
            item["source"] = "import"
        if not item.get("name"):
            errors.append(dict(row=row_index, message="이름 컬럼이 비어 있어 건너뜁니다."))
            continue
        if append_unmapped_to_memo:
            notes = []
            for index, value in enumerate(row):
                if mappings.get(index):
                    continue
                value = _cell_text(value)
                if not value:
                    continue
                header = column_headers.get(index) or f"열 {_column_letter(index)}"
                notes.append(f"{header}: {value}")
            if notes:
                extra = " / ".join(notes)
                item["memo"] = " / ".join([value for value in [item.get("memo", "").strip(), extra] if value])
        item["_row"] = row_index
        rows.append(item)

    return rows, errors


def _row_value(row, key):
    value = row.get(key, "")
    if isinstance(value, (datetime.datetime, datetime.date)):
        return _cell_text(value)
    return str(value or "")


def _export_csv(rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([label for _, label in EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([_row_value(row, key) for key, _ in EXPORT_COLUMNS])
    return output.getvalue().encode("utf-8-sig")


def _export_xlsx(rows):
    try:
        import openpyxl
    except Exception:
        raise Exception("XLSX 파일 생성을 위한 openpyxl 패키지가 필요합니다.")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Business Cards"
    sheet.append([label for _, label in EXPORT_COLUMNS])
    for row in rows:
        sheet.append([_row_value(row, key) for key, _ in EXPORT_COLUMNS])
    for column in sheet.columns:
        letter = column[0].column_letter
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 36)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def list():
    user_id = _current_user_id()
    text = wiz.request.query("text", "")
    scope = wiz.request.query("scope", "name_company")
    sort = wiz.request.query("sort", "name")
    direction = wiz.request.query("direction", "asc")
    page = int(wiz.request.query("page", 1))
    dump = int(wiz.request.query("dump", 10))
    rows, total = struct.card.search(user_id, text=text, scope=scope, sort=sort, direction=direction, page=page, dump=dump)
    all_total = struct.card.count(user_id)
    wiz.response.status(200, rows=rows, total=total, all_total=all_total)


def preview_import():
    _current_user_id()
    filename = wiz.request.query("filename", "")
    file_data = wiz.request.query("file_data", "")
    delimiter = wiz.request.query("delimiter", "auto")
    has_header = _bool_query("has_header", True)

    try:
        preview = _build_import_preview(filename, file_data, delimiter_mode=delimiter, has_header=has_header)
    except Exception as e:
        wiz.response.status(400, message=str(e))

    wiz.response.status(200, preview=preview)


def import_cards():
    user_id = _current_user_id()
    filename = wiz.request.query("filename", "")
    file_data = wiz.request.query("file_data", "")
    delimiter = wiz.request.query("delimiter", "auto")
    has_header = _bool_query("has_header", True)
    append_unmapped_to_memo = _bool_query("append_unmapped_to_memo", True)
    duplicate = wiz.request.query("duplicate", "skip")
    if duplicate not in ["skip", "update", "create"]:
        duplicate = "skip"

    try:
        preview = _build_import_preview(filename, file_data, delimiter_mode=delimiter, has_header=has_header)
        mapping = _load_mappings(wiz.request.query("mappings", "[]"), preview.get("mappings", []))
        mapping_errors = _mapping_duplicate_errors(mapping, preview.get("columns", []))
        if mapping_errors:
            messages = [
                f"{error.get('label')}: {', '.join(error.get('columns', []))}"
                for error in mapping_errors
            ]
            raise Exception("중복된 컬럼 매핑을 정리해주세요. " + " / ".join(messages))
        rows, errors = _mapped_import_rows(preview, mapping, append_unmapped_to_memo=append_unmapped_to_memo)
        result = struct.card.import_rows(user_id, rows, duplicate=duplicate)
        result["errors"] = errors + result.get("errors", [])
    except Exception as e:
        wiz.response.status(400, message=str(e))

    wiz.response.status(200, result=result)


def export_cards():
    user_id = _current_user_id()
    export_format = wiz.request.query("format", "xlsx")
    if export_format not in ["csv", "xlsx"]:
        export_format = "xlsx"

    text = wiz.request.query("text", "")
    scope = wiz.request.query("scope", "name_company")
    sort = wiz.request.query("sort", "name")
    direction = wiz.request.query("direction", "asc")

    try:
        rows, total = struct.card.export_rows(
            user_id,
            text=text,
            scope=scope,
            sort=sort,
            direction=direction,
            limit=MAX_EXPORT_ROWS,
        )
        if export_format == "csv":
            raw = _export_csv(rows)
            mime = "text/csv;charset=utf-8"
        else:
            raw = _export_xlsx(rows)
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        today = datetime.datetime.now().strftime("%Y%m%d")
        filename = f"{_safe_export_name('business-cards')}-{today}.{export_format}"
    except Exception as e:
        wiz.response.status(400, message=str(e))

    wiz.response.status(
        200,
        filename=filename,
        mime=mime,
        data=base64.b64encode(raw).decode("ascii"),
        count=len(rows),
        total=total,
        limited=total > len(rows),
    )


def analyze():
    _current_user_id()
    image_items = []
    front_image = wiz.request.query("front_image", "")
    back_image = wiz.request.query("back_image", "")

    if front_image:
        image_items.append(dict(side="front", label="앞면", image=front_image, filename=wiz.request.query("front_filename", "")))
    if back_image:
        image_items.append(dict(side="back", label="뒷면", image=back_image, filename=wiz.request.query("back_filename", "")))

    if not image_items:
        image_items.append(dict(
            side="single",
            label="명함",
            image=wiz.request.query("image", ""),
            filename=wiz.request.query("filename", ""),
        ))

    ai_enabled = False
    try:
        ai_enabled = struct.ai_setting.active() is not None
    except Exception:
        ai_enabled = False

    results = []
    for item in image_items:
        fallback_reason = ""
        image = None
        try:
            image = _decode_image(item["image"])
        except Exception as e:
            results.append(dict(
                side=item["side"],
                label=item["label"],
                filename=item["filename"],
                available=False,
                text="",
                confidence=0,
                fields={},
                message=str(e),
                elapsed_ms=0,
                passes=0,
                rotation=0,
                engine="image-invalid",
                fallback=False,
            ))
            continue

        ocr = None
        used_ai = False
        if ai_enabled:
            ocr = _ai_ocr(item["image"], item["label"])
            used_ai = ocr.get("available", False) and not ocr.get("execution_failed", False)
            if not used_ai:
                fallback_reason = ocr.get("message", "")

        if not used_ai:
            try:
                ocr = _server_ocr(image)
            except Exception as e:
                ocr = dict(available=False, text="", confidence=0, message=str(e), elapsed_ms=0, passes=0)

        fields = _merge_ai_fields(ocr.get("fields", {}), ocr.get("text", "")) if used_ai else _parse_ocr_text(ocr.get("text", ""))
        message = ocr.get("message", "")
        if fallback_reason and message:
            message = f"{message} (AI fallback: {fallback_reason})"
        results.append(dict(
            side=item["side"],
            label=item["label"],
            filename=item["filename"],
            available=ocr.get("available", False),
            text=ocr.get("text", ""),
            confidence=ocr.get("confidence", 0),
            fields=fields,
            message=message,
            elapsed_ms=ocr.get("elapsed_ms", 0),
            passes=ocr.get("passes", 0),
            rotation=ocr.get("rotation", 0),
            engine=ocr.get("engine", "ai" if used_ai else "server-tesseract"),
            engine_label=ocr.get("engine_label", ""),
            provider=ocr.get("provider", ""),
            model=ocr.get("model", ""),
            quality_notes=ocr.get("quality_notes", []),
            quality_passed=ocr.get("quality_passed", True),
            quality_score=ocr.get("quality_score", 0),
            payload_format=ocr.get("payload_format", "json" if used_ai else ""),
            execution_failed=ocr.get("execution_failed", False),
            attempts=ocr.get("attempts", 0),
            fallback=ai_enabled and not used_ai,
            fallback_reason=fallback_reason,
        ))

    available_results = [item for item in results if item.get("available")]
    merged_fields = _merge_ocr_fields([item.get("fields", {}) for item in results])
    combined_text = "\n\n".join([
        f"[{item.get('label')}]\n{item.get('text', '')}".strip()
        for item in results
        if item.get("text")
    ])
    confidence = 0
    if available_results:
        confidence = round(sum(item.get("confidence", 0) for item in available_results) / len(available_results))
    total_elapsed_ms = sum(item.get("elapsed_ms", 0) for item in results)
    total_passes = sum(item.get("passes", 0) for item in results)
    rotations = [item.get("rotation", 0) for item in results if item.get("rotation", 0)]
    engines = _unique_values([item.get("engine", "") for item in results if item.get("engine")])
    engine_labels = _unique_values([item.get("engine_label", "") for item in results if item.get("engine_label")])
    used_ai = any(str(engine).startswith("ai-") for engine in engines)
    used_server = any(str(engine).startswith("server-") for engine in engines)
    fallback_used = ai_enabled and used_server
    ai_quality_failed = any(
        str(item.get("engine", "")).startswith("ai-") and item.get("quality_passed") is False
        for item in results
        if item.get("available")
    )
    if used_ai and fallback_used:
        message = "AI 분석 완료, 일부 서버 fallback"
    elif used_ai and ai_quality_failed:
        message = "AI 분석 완료, 품질 확인 필요"
    elif used_ai:
        message = "AI 분석 완료"
    elif fallback_used and len(available_results) > 0:
        message = "서버 fallback 분석 완료"
    elif len(available_results) > 0:
        message = "서버 분석 완료"
    else:
        message = "서버 OCR 미가용"

    wiz.response.status(
        200,
        available=len(available_results) > 0,
        engine=", ".join(engines) if engines else ("server-tesseract" if len(available_results) > 0 else "server-unavailable"),
        engine_label=", ".join(engine_labels) if engine_labels else "",
        filename=", ".join([item.get("filename", "") for item in results if item.get("filename")]),
        text=combined_text,
        confidence=confidence,
        fields=merged_fields,
        results=results,
        message=message,
        ai_enabled=ai_enabled,
        quality_passed=not ai_quality_failed,
        fallback=fallback_used or len(available_results) == 0,
        elapsed_ms=total_elapsed_ms,
        passes=total_passes,
        rotations=rotations,
    )


def get():
    user_id = _current_user_id()
    card_id = wiz.request.query("id", "")
    if not card_id:
        wiz.response.status(400, message="명함 ID가 필요합니다.")
    card = struct.card.get(card_id, user_id)
    if card is None:
        wiz.response.status(404, message="명함을 찾을 수 없습니다.")
    wiz.response.status(200, card=card)


def save():
    user_id = _current_user_id()
    card_id = wiz.request.query("id", "")
    data = _payload()

    try:
        if card_id:
            struct.card.update(card_id, user_id, data)
        else:
            card_id = struct.card.create(user_id, data)
    except Exception as e:
        wiz.response.status(400, message=str(e))

    card = struct.card.get(card_id, user_id)
    wiz.response.status(200, card=card)


def remove():
    user_id = _current_user_id()
    card_id = wiz.request.query("id", "")
    if not card_id:
        wiz.response.status(400, message="명함 ID가 필요합니다.")

    try:
        struct.card.delete(card_id, user_id)
    except Exception as e:
        wiz.response.status(400, message=str(e))

    wiz.response.status(200)
