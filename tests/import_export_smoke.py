import base64
import importlib.util
import io
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
PROJECT = ROOT / "project" / "main"
API_PATH = PROJECT / "src" / "app" / "page.cards" / "api.py"


class FakeWiz:
    def model(self, _name):
        return None


def load_api():
    spec = importlib.util.spec_from_file_location("cards_api", API_PATH)
    module = importlib.util.module_from_spec(spec)
    module.wiz = FakeWiz()
    spec.loader.exec_module(module)
    return module


def data_url(raw, mime="text/csv"):
    return f"data:{mime};base64,{base64.b64encode(raw).decode('ascii')}"


def mapping_dict(preview):
    return {
        int(item["index"]): item["field"]
        for item in preview["mappings"]
        if item.get("field")
    }


def assert_csv_import(api):
    csv_text = (
        "이름,회사,부서,직책,이메일,휴대폰,전화,주소,웹사이트,메모\n"
        "홍길동,나노하테크,개발팀,팀장,hong@example.com,010-1111-2222,02-123-4567,서울,example.com,VIP\n"
    )
    preview = api._build_import_preview("cards.csv", data_url(csv_text.encode("utf-8-sig")), has_header=True)
    fields = [item["field"] for item in preview["mappings"]]
    for expected in ["name", "company", "department", "position", "email", "mobile", "phone", "address", "website", "memo"]:
        if expected not in fields:
            raise AssertionError(f"CSV header was not mapped to {expected}")

    rows, errors = api._mapped_import_rows(preview, mapping_dict(preview))
    if errors:
        raise AssertionError(f"unexpected CSV import errors: {errors}")
    if rows[0]["name"] != "홍길동" or rows[0]["source"] != "import":
        raise AssertionError(f"unexpected CSV row mapping: {rows[0]}")


def assert_mapping_validation(api):
    csv_text = (
        "이름,명함첩 이름,추가정보,이메일\n"
        "홍길동,거래처,VIP,hong@example.com\n"
    )
    preview = api._build_import_preview("cards.csv", data_url(csv_text.encode("utf-8")), has_header=True)
    mapping = mapping_dict(preview)
    errors = api._mapping_duplicate_errors(mapping, preview["columns"])
    if not errors or errors[0]["field"] != "name":
        raise AssertionError(f"expected duplicate name mapping error: {errors}")

    mapping[1] = ""
    rows, row_errors = api._mapped_import_rows(preview, mapping, append_unmapped_to_memo=True)
    if row_errors:
        raise AssertionError(f"unexpected row errors after duplicate fix: {row_errors}")
    if "명함첩 이름: 거래처 / 추가정보: VIP" not in rows[0]["memo"]:
        raise AssertionError(f"unmapped column was not appended to memo: {rows[0]}")


def assert_xlsx_import(api):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Company", "Email", "Mobile"])
    sheet.append(["Jane Doe", "Nanoha Labs", "jane@example.com", "+1 415 555 1234"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    preview = api._build_import_preview(
        "cards.xlsx",
        data_url(output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        has_header=True,
    )
    rows, errors = api._mapped_import_rows(preview, mapping_dict(preview))
    if errors:
        raise AssertionError(f"unexpected XLSX import errors: {errors}")
    if rows[0]["name"] != "Jane Doe" or rows[0]["company"] != "Nanoha Labs":
        raise AssertionError(f"unexpected XLSX row mapping: {rows[0]}")


def assert_export(api):
    rows = [
        {
            "name": "홍길동",
            "company": "나노하테크",
            "email": "hong@example.com",
            "mobile": "010-1111-2222",
        }
    ]
    csv_bytes = api._export_csv(rows)
    if "홍길동" not in csv_bytes.decode("utf-8-sig"):
        raise AssertionError("CSV export missing expected row")
    xlsx_bytes = api._export_xlsx(rows)
    loaded = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    values = list(loaded.active.iter_rows(values_only=True))
    loaded.close()
    if values[1][0] != "홍길동":
        raise AssertionError(f"XLSX export missing expected row: {values}")


def main():
    api = load_api()
    assert_csv_import(api)
    assert_mapping_validation(api)
    assert_xlsx_import(api)
    assert_export(api)
    print("import/export helpers: ok")


if __name__ == "__main__":
    main()
